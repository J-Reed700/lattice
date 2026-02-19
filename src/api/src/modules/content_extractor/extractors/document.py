from datetime import datetime
from pathlib import Path
from typing import Any

from ..detector import check_file_size, detect_mime_type
from ..types import CorruptedFileError, ExtractedContent, ExtractionError
from ..utils import detect_language, normalize_text
from . import register_extractor


@register_extractor("application/vnd.openxmlformats-officedocument.wordprocessingml.document")
def extract_docx(file_path: str) -> ExtractedContent:
    """Extract content from DOCX files using python-docx.

    Extracts text from paragraphs and tables.
    Extracts metadata including author, title, etc.

    Args:
        file_path: Path to the DOCX file

    Returns:
        ExtractedContent with extracted text and document metadata

    Raises:
        CorruptedFileError: If DOCX structure is invalid
        ExtractionError: If extraction fails
        FileSizeLimitError: If file exceeds size limit

    Example:
        >>> content = extract_docx('/path/to/document.docx')
        >>> assert 'application/vnd.openxmlformats' in content.mime_type
        >>> assert 'paragraph_count' in content.metadata
    """
    try:
        from docx import Document
    except ImportError as e:
        raise ExtractionError(
            "python-docx library not installed. Install with: pip install python-docx",
            file_path=file_path,
        ) from e

    try:
        mime_type = detect_mime_type(file_path)
        file_size = check_file_size(file_path, mime_type)

        try:
            doc = Document(file_path)
        except Exception as e:
            raise CorruptedFileError(f"Failed to open DOCX file: {e!s}", file_path=file_path) from e

        text_parts = []
        for paragraph in doc.paragraphs:
            if paragraph.text.strip():
                text_parts.append(paragraph.text)

        for table in doc.tables:
            for row in table.rows:
                row_text = []
                for cell in row.cells:
                    if cell.text.strip():
                        row_text.append(cell.text.strip())
                if row_text:
                    text_parts.append(" | ".join(row_text))

        text = "\n\n".join(text_parts)
        text = normalize_text(text)

        metadata = _extract_docx_metadata(doc)
        metadata["paragraph_count"] = len(doc.paragraphs)
        metadata["table_count"] = len(doc.tables)
        metadata["file_name"] = Path(file_path).name
        metadata["file_path"] = str(Path(file_path).absolute())

        language = detect_language(text)

        return ExtractedContent(
            text=text,
            mime_type=mime_type,
            file_path=file_path,
            metadata=metadata,
            encoding=None,
            language=language,
            extraction_time=datetime.now(),
            file_size=file_size,
        )

    except (CorruptedFileError, ExtractionError):
        raise
    except Exception as e:
        raise ExtractionError(f"Failed to extract DOCX content: {e!s}", file_path=file_path) from e


@register_extractor("application/vnd.openxmlformats-officedocument.presentationml.presentation")
def extract_pptx(file_path: str) -> ExtractedContent:
    """Extract content from PPTX files using python-pptx.

    Extracts text from all slides and shapes.
    Includes slide count and presentation metadata.

    Args:
        file_path: Path to the PPTX file

    Returns:
        ExtractedContent with extracted text and presentation metadata

    Raises:
        CorruptedFileError: If PPTX structure is invalid
        ExtractionError: If extraction fails
        FileSizeLimitError: If file exceeds size limit

    Example:
        >>> content = extract_pptx('/path/to/presentation.pptx')
        >>> assert 'slide_count' in content.metadata
    """
    try:
        from pptx import Presentation
    except ImportError as e:
        raise ExtractionError(
            "python-pptx library not installed. Install with: pip install python-pptx",
            file_path=file_path,
        ) from e

    try:
        mime_type = detect_mime_type(file_path)
        file_size = check_file_size(file_path, mime_type)

        try:
            prs = Presentation(file_path)
        except Exception as e:
            raise CorruptedFileError(f"Failed to open PPTX file: {e!s}", file_path=file_path) from e

        text_parts = []
        for slide_num, slide in enumerate(prs.slides, 1):
            slide_text = []
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text.strip():
                    slide_text.append(shape.text)

            if slide_text:
                text_parts.append(f"[Slide {slide_num}]\n" + "\n".join(slide_text))

        text = "\n\n".join(text_parts)
        text = normalize_text(text)

        metadata = _extract_pptx_metadata(prs)
        metadata["slide_count"] = len(prs.slides)
        metadata["file_name"] = Path(file_path).name
        metadata["file_path"] = str(Path(file_path).absolute())

        language = detect_language(text)

        return ExtractedContent(
            text=text,
            mime_type=mime_type,
            file_path=file_path,
            metadata=metadata,
            encoding=None,
            language=language,
            extraction_time=datetime.now(),
            file_size=file_size,
        )

    except (CorruptedFileError, ExtractionError):
        raise
    except Exception as e:
        raise ExtractionError(f"Failed to extract PPTX content: {e!s}", file_path=file_path) from e


@register_extractor("application/vnd.oasis.opendocument.text")
def extract_odt(file_path: str) -> ExtractedContent:
    """Extract content from ODT files using odfpy.

    Extracts text from OpenDocument Text format.

    Args:
        file_path: Path to the ODT file

    Returns:
        ExtractedContent with extracted text

    Raises:
        CorruptedFileError: If ODT structure is invalid
        ExtractionError: If extraction fails
        FileSizeLimitError: If file exceeds size limit

    Example:
        >>> content = extract_odt('/path/to/document.odt')
        >>> assert content.mime_type == 'application/vnd.oasis.opendocument.text'
    """
    try:
        from odf import text as odf_text
        from odf.opendocument import load
    except ImportError as e:
        raise ExtractionError(
            "odfpy library not installed. Install with: pip install odfpy", file_path=file_path
        ) from e

    try:
        mime_type = detect_mime_type(file_path)
        file_size = check_file_size(file_path, mime_type)

        try:
            doc = load(file_path)
        except Exception as e:
            raise CorruptedFileError(f"Failed to open ODT file: {e!s}", file_path=file_path) from e

        text_parts = []
        for paragraph in doc.getElementsByType(odf_text.P):
            para_text = str(paragraph)
            if para_text.strip():
                text_parts.append(para_text)

        text = "\n\n".join(text_parts)
        text = normalize_text(text)

        metadata = {}
        metadata["file_name"] = Path(file_path).name
        metadata["file_path"] = str(Path(file_path).absolute())

        language = detect_language(text)

        return ExtractedContent(
            text=text,
            mime_type=mime_type,
            file_path=file_path,
            metadata=metadata,
            encoding=None,
            language=language,
            extraction_time=datetime.now(),
            file_size=file_size,
        )

    except (CorruptedFileError, ExtractionError):
        raise
    except Exception as e:
        raise ExtractionError(f"Failed to extract ODT content: {e!s}", file_path=file_path) from e


def _extract_docx_metadata(doc) -> dict[str, Any]:
    """Extract metadata from DOCX document.

    Args:
        doc: python-docx Document object

    Returns:
        Dictionary with metadata fields
    """
    metadata = {}

    try:
        core_props = doc.core_properties
        if core_props.author:
            metadata["author"] = core_props.author
        if core_props.title:
            metadata["title"] = core_props.title
        if core_props.subject:
            metadata["subject"] = core_props.subject
        if core_props.keywords:
            metadata["keywords"] = core_props.keywords
        if core_props.created:
            metadata["created"] = core_props.created.isoformat()
        if core_props.modified:
            metadata["modified"] = core_props.modified.isoformat()
    except Exception:
        pass

    return metadata


def _extract_pptx_metadata(prs) -> dict[str, Any]:
    """Extract metadata from PPTX presentation.

    Args:
        prs: python-pptx Presentation object

    Returns:
        Dictionary with metadata fields
    """
    metadata = {}

    try:
        core_props = prs.core_properties
        if core_props.author:
            metadata["author"] = core_props.author
        if core_props.title:
            metadata["title"] = core_props.title
        if core_props.subject:
            metadata["subject"] = core_props.subject
        if core_props.created:
            metadata["created"] = core_props.created.isoformat()
        if core_props.modified:
            metadata["modified"] = core_props.modified.isoformat()
    except Exception:
        pass

    return metadata
